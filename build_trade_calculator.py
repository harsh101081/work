"""
Generate a professional Forex & Prop-Firm trade calculator workbook.

The user only types the Entry Price (plus one-time settings); the sheet
calculates Stop-Loss, Take-Profit, lot size, dollar risk and reward using
live Excel formulas. Also includes a multi-row trade planner, an instrument
spec lookup table and an optional prop-firm rule check.

Run:  python3 build_trade_calculator.py
Out:  Forex_Trade_Calculator.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Forex_Trade_Calculator.xlsx"

# ---- Palette -------------------------------------------------------------
NAVY = "0B1220"
SURFACE = "131C2E"
ACCENT = "1E3A8A"
ACCENT_BRIGHT = "2563EB"
SUCCESS = "16A34A"
DANGER = "DC2626"
WARNING = "B45309"
INPUT_FILL = "FEF3C7"      # soft yellow -> "edit me"
ENTRY_FILL = "BBF7D0"      # green -> the single per-trade input
HEADER_FILL = "1E3A8A"
ZEBRA = "F1F5F9"
GRID = "CBD5E1"

WHITE = "FFFFFF"
MUTED = "64748B"
INK = "0F172A"

# ---- Number formats ------------------------------------------------------
FMT_MONEY = '$#,##0.00'
FMT_PRICE = '#,##0.00####'
FMT_LOT = '0.00'
FMT_PIPS = '0.0'
FMT_PIPSIZE = '0.00####'
FMT_PCT = '0.00%'
FMT_INT = '0'

# ---- Shared style objects ------------------------------------------------
thin = Side(style="thin", color=GRID)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(color=INK, bold=False, size=11, italic=False):
    return Font(name="Calibri", color=color, bold=bold, size=size, italic=italic)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def label(ws, cell, text, bold=False, color=INK):
    c = ws[cell]
    c.value = text
    c.font = font(color=color, bold=bold)
    c.alignment = Alignment(horizontal="left", vertical="center")
    return c


def value_cell(ws, cell, value, *, number_format=None, fill_hex=None,
               color=INK, bold=False, border=True, align="right", size=11):
    c = ws[cell]
    c.value = value
    c.font = font(color=color, bold=bold, size=size)
    if number_format:
        c.number_format = number_format
    if fill_hex:
        c.fill = fill(fill_hex)
    if border:
        c.border = border_all
    c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def section(ws, row, text, last_col="D"):
    ws.merge_cells(f"B{row}:{last_col}{row}")
    c = ws[f"B{row}"]
    c.value = text
    c.fill = fill(HEADER_FILL)
    c.font = font(color=WHITE, bold=True, size=11)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 22


def note(ws, cell, text):
    c = ws[cell]
    c.value = text
    c.font = font(color=MUTED, size=9, italic=True)
    c.alignment = Alignment(horizontal="left", vertical="center")
    return c


# =========================================================================
#  Workbook
# =========================================================================
wb = Workbook()

# -------------------------------------------------------------------------
#  Sheet 1: Calculator
# -------------------------------------------------------------------------
ws = wb.active
ws.title = "Calculator"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = ACCENT_BRIGHT

ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 30
ws.column_dimensions["E"].width = 2

# Title
ws.merge_cells("B2:D2")
t = ws["B2"]
t.value = "FOREX  &  PROP-FIRM  TRADE  CALCULATOR"
t.fill = fill(NAVY)
t.font = font(color=WHITE, bold=True, size=16)
t.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 34

ws.merge_cells("B3:D3")
s = ws["B3"]
s.value = "Set your inputs once, then for each trade just type the green Entry Price cell."
s.font = font(color=MUTED, italic=True, size=10)
s.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[3].height = 18

# ---- Section 1: Settings
section(ws, 5, "1  -  SETTINGS  (edit once)")
label(ws, "B6", "Account Balance ($)")
value_cell(ws, "C6", 2488.76, number_format=FMT_MONEY, fill_hex=INPUT_FILL, bold=True)
note(ws, "D6", "<- your current balance used for risk sizing")

label(ws, "B7", "Risk % per Trade")
value_cell(ws, "C7", 0.01, number_format=FMT_PCT, fill_hex=INPUT_FILL, bold=True)
note(ws, "D7", "<- e.g. 1.00%  (keep <= 1-2%)")

label(ws, "B8", "Instrument")
value_cell(ws, "C8", "XAUUSD", fill_hex=INPUT_FILL, bold=True, align="center")
note(ws, "D8", "<- dropdown: EURUSD / GBPUSD / XAUUSD")

label(ws, "B9", "Direction")
value_cell(ws, "C9", "BUY", fill_hex=INPUT_FILL, bold=True, align="center")
note(ws, "D9", "<- dropdown: BUY / SELL")

label(ws, "B10", "Stop-Loss Distance (pips)")
value_cell(ws, "C10", 25, number_format=FMT_PIPS, fill_hex=INPUT_FILL, bold=True)
note(ws, "D10", "<- defines your risk distance")

label(ws, "B11", "Risk : Reward  (1 : x)")
value_cell(ws, "C11", 2, number_format='0.0#', fill_hex=INPUT_FILL, bold=True)
note(ws, "D11", "<- 1, 1.5, 2, 3, 4, 5 or custom")

# ---- Section 2: Entry (the only per-trade input)
section(ws, 13, "2  -  ENTRY  (the only thing you change per trade)")
label(ws, "B14", "Entry Price", bold=True)
value_cell(ws, "C14", 2350.0, number_format=FMT_PRICE, fill_hex=ENTRY_FILL,
           bold=True, size=13)
note(ws, "D14", "<- TYPE YOUR ENTRY HERE")
ws.row_dimensions[14].height = 24

# ---- Section 3: Instrument data (auto via VLOOKUP)
LK = "Instruments!$A$2:$F$4"
section(ws, 16, "3  -  INSTRUMENT DATA  (auto)")
label(ws, "B17", "Pip Size")
value_cell(ws, "C17", f"=VLOOKUP($C$8,{LK},2,FALSE)", number_format=FMT_PIPSIZE)
label(ws, "B18", "Pip Value per Lot ($)")
value_cell(ws, "C18", f"=VLOOKUP($C$8,{LK},3,FALSE)", number_format=FMT_MONEY)
label(ws, "B19", "Price Decimals")
value_cell(ws, "C19", f"=VLOOKUP($C$8,{LK},5,FALSE)", number_format=FMT_INT)

# ---- Section 4: Results (auto)
section(ws, 21, "4  -  RESULTS  (auto-calculated)")

label(ws, "B22", "Stop-Loss Price", bold=True)
value_cell(ws, "C22",
           "=IF($C$9=\"BUY\",$C$14-$C$10*$C$17,$C$14+$C$10*$C$17)",
           number_format=FMT_PRICE, color=DANGER, bold=True)

label(ws, "B23", "Take-Profit Price", bold=True)
value_cell(ws, "C23",
           "=IF($C$9=\"BUY\",$C$14+$C$10*$C$17*$C$11,$C$14-$C$10*$C$17*$C$11)",
           number_format=FMT_PRICE, color=SUCCESS, bold=True)

label(ws, "B24", "Stop Distance (price)")
value_cell(ws, "C24", "=$C$10*$C$17", number_format=FMT_PRICE)

label(ws, "B25", "Reward Distance (pips)")
value_cell(ws, "C25", "=$C$10*$C$11", number_format=FMT_PIPS)

label(ws, "B26", "Dollar Risk ($)")
value_cell(ws, "C26", "=$C$6*$C$7", number_format=FMT_MONEY, color=WARNING, bold=True)

label(ws, "B27", "Recommended Lot (raw)")
value_cell(ws, "C27", "=IF($C$10*$C$18=0,0,$C$26/($C$10*$C$18))",
           number_format='0.000')

label(ws, "B28", "Recommended Lot Size", bold=True)
value_cell(ws, "C28", "=ROUND($C$27,2)", number_format=FMT_LOT,
           color=ACCENT_BRIGHT, bold=True, size=12)

label(ws, "B29", "Actual Risk at Lot ($)")
value_cell(ws, "C29", "=$C$28*$C$10*$C$18", number_format=FMT_MONEY, color=DANGER)

label(ws, "B30", "Potential Profit ($)", bold=True)
value_cell(ws, "C30", "=$C$29*$C$11", number_format=FMT_MONEY, color=SUCCESS, bold=True)

label(ws, "B31", "Risk : Reward")
value_cell(ws, "C31", "=\"1 : \"&TEXT($C$11,\"0.##\")", align="center")

# ---- Section 5: Prop-firm check
section(ws, 33, "5  -  PROP-FIRM RULE CHECK  (optional)")
label(ws, "B34", "Starting Balance ($)")
value_cell(ws, "C34", 2500, number_format=FMT_MONEY, fill_hex=INPUT_FILL)
label(ws, "B35", "Current Equity ($)")
value_cell(ws, "C35", 2488.76, number_format=FMT_MONEY, fill_hex=INPUT_FILL)
label(ws, "B36", "Daily Drawdown Limit (%)")
value_cell(ws, "C36", 0.03, number_format=FMT_PCT, fill_hex=INPUT_FILL)
label(ws, "B37", "Max Drawdown Limit (%)")
value_cell(ws, "C37", 0.06, number_format=FMT_PCT, fill_hex=INPUT_FILL)
label(ws, "B38", "Profit Target (%)")
value_cell(ws, "C38", 0.03, number_format=FMT_PCT, fill_hex=INPUT_FILL)

label(ws, "B39", "Remaining Daily Drawdown ($)")
value_cell(ws, "C39", "=MAX(0,$C$35-($C$6-$C$34*$C$36))",
           number_format=FMT_MONEY, color=WARNING)
label(ws, "B40", "Remaining Max Drawdown ($)")
value_cell(ws, "C40", "=MAX(0,$C$35-($C$34*(1-$C$37)))",
           number_format=FMT_MONEY, color=DANGER)
label(ws, "B41", "Losing Trades left (Daily)")
value_cell(ws, "C41", "=IF($C$29<=0,0,FLOOR($C$39/$C$29,1))", number_format=FMT_INT)
label(ws, "B42", "Losing Trades left (Max)")
value_cell(ws, "C42", "=IF($C$29<=0,0,FLOOR($C$40/$C$29,1))", number_format=FMT_INT)
label(ws, "B43", "Losing Trades remaining", bold=True)
value_cell(ws, "C43", "=MIN($C$41,$C$42)", number_format=FMT_INT,
           color=DANGER, bold=True)

label(ws, "B44", "Profit Target ($)")
value_cell(ws, "C44", "=$C$34*$C$38", number_format=FMT_MONEY)
label(ws, "B45", "Amount to Reach Target ($)")
value_cell(ws, "C45", "=MAX(0,($C$34+$C$44)-$C$6)",
           number_format=FMT_MONEY, color=SUCCESS)
label(ws, "B46", "Winning Trades to Pass", bold=True)
value_cell(ws, "C46", "=IF($C$30<=0,\"-\",ROUNDUP($C$45/$C$30,0))",
           number_format=FMT_INT, color=SUCCESS, bold=True)

note(ws, "B48", "Formulas recalc instantly. This tool is for planning only -"
                " confirm pip values & rules with your broker / prop firm.")

# ---- Data validation (dropdowns) on Calculator
dv_instr = DataValidation(type="list", formula1='"EURUSD,GBPUSD,XAUUSD"',
                          allow_blank=False)
dv_dir = DataValidation(type="list", formula1='"BUY,SELL"', allow_blank=False)
dv_rr = DataValidation(type="list", formula1='"1,1.5,2,3,4,5"',
                       allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv_instr)
ws.add_data_validation(dv_dir)
ws.add_data_validation(dv_rr)
dv_instr.add(ws["C8"])
dv_dir.add(ws["C9"])
dv_rr.add(ws["C11"])

# -------------------------------------------------------------------------
#  Sheet 2: Trade Planner (multi-row)
# -------------------------------------------------------------------------
pw = wb.create_sheet("Trade Planner")
pw.sheet_view.showGridLines = False
pw.sheet_properties.tabColor = SUCCESS

widths = {"A": 2, "B": 5, "C": 11, "D": 14, "E": 14, "F": 14,
          "G": 14, "H": 11, "I": 16, "J": 20}
for col, w in widths.items():
    pw.column_dimensions[col].width = w

pw.merge_cells("B2:J2")
pt = pw["B2"]
pt.value = "TRADE PLANNER  -  enter Direction + Entry Price on each row"
pt.fill = fill(NAVY)
pt.font = font(color=WHITE, bold=True, size=14)
pt.alignment = Alignment(horizontal="center", vertical="center")
pw.row_dimensions[2].height = 30

# Settings block
section(pw, 4, "SETTINGS", last_col="C")
label(pw, "B5", "Account Balance ($)")
value_cell(pw, "C5", 2488.76, number_format=FMT_MONEY, fill_hex=INPUT_FILL, bold=True)
label(pw, "B6", "Risk % per Trade")
value_cell(pw, "C6", 0.01, number_format=FMT_PCT, fill_hex=INPUT_FILL, bold=True)
label(pw, "B7", "Instrument")
value_cell(pw, "C7", "XAUUSD", fill_hex=INPUT_FILL, bold=True, align="center")
label(pw, "B8", "Stop-Loss (pips)")
value_cell(pw, "C8", 25, number_format=FMT_PIPS, fill_hex=INPUT_FILL, bold=True)
label(pw, "B9", "Risk : Reward (1 : x)")
value_cell(pw, "C9", 2, number_format='0.0#', fill_hex=INPUT_FILL, bold=True)

# Derived helper cells (col E/F)
label(pw, "E5", "Pip Size:")
value_cell(pw, "F5", f"=VLOOKUP($C$7,{LK},2,FALSE)", number_format=FMT_PIPSIZE)
label(pw, "E6", "Pip Value / Lot:")
value_cell(pw, "F6", f"=VLOOKUP($C$7,{LK},3,FALSE)", number_format=FMT_MONEY)
label(pw, "E7", "Dollar Risk / Trade:")
value_cell(pw, "F7", "=$C$5*$C$6", number_format=FMT_MONEY, color=WARNING, bold=True)

# Table header
HEAD_ROW = 11
headers = ["#", "Direction", "Entry Price", "Stop-Loss", "Take-Profit",
           "Dollar Risk", "Lot Size", "Potential Profit", "Notes"]
for i, h in enumerate(headers):
    col = chr(ord("B") + i)
    c = pw[f"{col}{HEAD_ROW}"]
    c.value = h
    c.fill = fill(HEADER_FILL)
    c.font = font(color=WHITE, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_all
pw.row_dimensions[HEAD_ROW].height = 20

FIRST = HEAD_ROW + 1
LAST = FIRST + 14  # 15 planner rows
for idx, r in enumerate(range(FIRST, LAST + 1)):
    zebra = ZEBRA if idx % 2 else WHITE
    # # column
    value_cell(pw, f"B{r}", idx + 1, number_format=FMT_INT,
               fill_hex=zebra, align="center")
    # Direction (input, default BUY)
    value_cell(pw, f"C{r}", "BUY", fill_hex=INPUT_FILL, align="center")
    # Entry price (input)
    value_cell(pw, f"D{r}", None, number_format=FMT_PRICE, fill_hex=ENTRY_FILL)
    # Stop-Loss
    value_cell(pw, f"E{r}",
               f'=IF($D{r}="","",IF($C{r}="SELL",$D{r}+$C$8*$F$5,$D{r}-$C$8*$F$5))',
               number_format=FMT_PRICE, color=DANGER, fill_hex=zebra)
    # Take-Profit
    value_cell(pw, f"F{r}",
               f'=IF($D{r}="","",IF($C{r}="SELL",$D{r}-$C$8*$F$5*$C$9,$D{r}+$C$8*$F$5*$C$9))',
               number_format=FMT_PRICE, color=SUCCESS, fill_hex=zebra)
    # Dollar risk
    value_cell(pw, f"G{r}", f'=IF($D{r}="","",$F$7)',
               number_format=FMT_MONEY, fill_hex=zebra)
    # Lot size
    value_cell(pw, f"H{r}",
               f'=IF($D{r}="","",ROUND($F$7/($C$8*$F$6),2))',
               number_format=FMT_LOT, color=ACCENT_BRIGHT, bold=True, fill_hex=zebra)
    # Potential profit
    value_cell(pw, f"I{r}", f'=IF($D{r}="","",$F$7*$C$9)',
               number_format=FMT_MONEY, color=SUCCESS, fill_hex=zebra)
    # Notes
    value_cell(pw, f"J{r}", None, fill_hex=zebra, align="left")

# Sample first row so the sheet isn't empty
pw["D12"] = 2350.0

note(pw, f"B{LAST + 2}", "Tip: copy a row down to add more trades. Change the"
                         " SETTINGS block to re-size every row at once.")

# Planner dropdowns
pdv_instr = DataValidation(type="list", formula1='"EURUSD,GBPUSD,XAUUSD"',
                           allow_blank=False)
pdv_dir = DataValidation(type="list", formula1='"BUY,SELL"', allow_blank=False)
pdv_rr = DataValidation(type="list", formula1='"1,1.5,2,3,4,5"',
                        allow_blank=True, showErrorMessage=False)
pw.add_data_validation(pdv_instr)
pw.add_data_validation(pdv_dir)
pw.add_data_validation(pdv_rr)
pdv_instr.add(pw["C7"])
pdv_rr.add(pw["C9"])
pdv_dir.add(f"C{FIRST}:C{LAST}")

# -------------------------------------------------------------------------
#  Sheet 3: Instruments (lookup table)
# -------------------------------------------------------------------------
iw = wb.create_sheet("Instruments")
iw.sheet_properties.tabColor = MUTED
iw.sheet_view.showGridLines = False

inst_headers = ["Symbol", "Pip Size", "Pip Value / Lot ($)",
                "Contract Size", "Price Decimals", "Reference Price"]
inst_rows = [
    ["EURUSD", 0.0001, 10, 100000, 5, 1.085],
    ["GBPUSD", 0.0001, 10, 100000, 5, 1.270],
    ["XAUUSD", 0.1, 10, 100, 2, 2350.0],
]
for i, h in enumerate(inst_headers):
    c = iw.cell(row=1, column=i + 1, value=h)
    c.fill = fill(HEADER_FILL)
    c.font = font(color=WHITE, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_all
for r, row in enumerate(inst_rows, start=2):
    for cidx, val in enumerate(row, start=1):
        c = iw.cell(row=r, column=cidx, value=val)
        c.border = border_all
        c.alignment = Alignment(horizontal="center", vertical="center")
        if cidx == 2:
            c.number_format = FMT_PIPSIZE
        elif cidx == 3 or cidx == 6:
            c.number_format = FMT_PRICE if cidx == 6 else FMT_MONEY
        elif cidx in (4, 5):
            c.number_format = FMT_INT
for col, w in {"A": 12, "B": 11, "C": 18, "D": 14, "E": 14, "F": 16}.items():
    iw.column_dimensions[col].width = w

iw["A6"] = ("Note: For these USD-quoted pairs, 1 pip per standard lot = $10. "
            "Gold (XAUUSD) assumes a 100-oz contract where a $0.10 move = $10/lot. "
            "Adjust these values to match your broker if needed.")
iw["A6"].font = font(color=MUTED, italic=True, size=9)

wb.save(OUT)
print("Saved", OUT)
